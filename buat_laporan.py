# -*- coding: utf-8 -*-
"""Generator laporan tugas akhir mata kuliah: Aplikasi Bangun Kelola."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ---------- Pengaturan gaya dasar ----------
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
pf.space_after = Pt(0)


def set_cell_font(cell, bold=False, size=11, white=False):
    for p in cell.paragraphs:
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(size)
            r.font.bold = bold
            if white:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def shade_cell(cell, color="1F4E79"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def H1(text):
    p = doc.add_heading(level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    return p


def H2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x2E, 0x5C, 0x8A)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def H3(text):
    p = doc.add_heading(level=3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    return p


def para(text, justify=True, italic=False, indent=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.4)
    r = p.add_run(text)
    r.font.italic = italic
    return p


def bullet(text, num=False):
    style = "List Number" if num else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    return p


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.italic = True
    return p


def figure_placeholder(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("[ " + text + " ]")
    r.font.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    return p


def table(headers, rows, widths=None, header_color="1F4E79"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        shade_cell(hdr[i], header_color)
        set_cell_font(hdr[i], bold=True, size=10.5, white=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            set_cell_font(cells[i], size=10.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def page_break():
    doc.add_page_break()


# ============================================================
# HALAMAN JUDUL
# ============================================================
def center_run(text, size=14, bold=True, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    r = p.add_run(text)
    r.font.bold = bold
    r.font.size = Pt(size)
    return p


center_run("LAPORAN PROYEK APLIKASI", 14, True, 4, 24)
center_run("BANGUN KELOLA: APLIKASI PEMBUKUAN DAN AKUNTANSI", 16, True, 2)
center_run("BERBASIS WEB UNTUK TOKO BANGUNAN", 16, True, 24)

figure_placeholder("Logo Institusi / Perguruan Tinggi")
doc.add_paragraph().paragraph_format.space_after = Pt(24)

center_run("Diajukan untuk Memenuhi Tugas Akhir Mata Kuliah", 12, False, 2, 12)
center_run("Pengembangan Aplikasi Sistem Informasi", 12, False, 24)

center_run("Disusun oleh:", 12, False, 2)
center_run("[Nama Penyusun]", 12, True, 2)
center_run("[Nomor Induk Mahasiswa]", 12, False, 24)

center_run("[NAMA PROGRAM STUDI]", 12, True, 2, 36)
center_run("[NAMA FAKULTAS]", 12, True, 2)
center_run("[NAMA PERGURUAN TINGGI]", 12, True, 2)
center_run("2026", 12, True, 2)

page_break()

# ============================================================
# KATA PENGANTAR
# ============================================================
H1("KATA PENGANTAR")
para("Puji syukur penyusun panjatkan ke hadirat Tuhan Yang Maha Esa atas limpahan rahmat dan karunia-Nya sehingga laporan proyek aplikasi yang berjudul “Bangun Kelola: Aplikasi Pembukuan dan Akuntansi Berbasis Web untuk Toko Bangunan” ini dapat diselesaikan dengan baik. Laporan ini disusun sebagai bentuk pertanggungjawaban akademik atas proyek pengembangan aplikasi yang menjadi tugas akhir mata kuliah.")
para("Laporan ini memuat uraian menyeluruh mulai dari latar belakang permasalahan yang dihadapi pelaku usaha toko bangunan, kajian teori yang melandasi pembangunan sistem, proses analisis dan perancangan, hingga implementasi serta pengujian aplikasi. Melalui laporan ini diharapkan pembaca dapat memahami bagaimana sebuah kebutuhan nyata di lapangan diterjemahkan menjadi sebuah solusi perangkat lunak yang utuh dan bermanfaat.")
para("Penyusun menyadari bahwa penyelesaian laporan ini tidak terlepas dari bimbingan dan dukungan berbagai pihak. Oleh karena itu, penyusun menyampaikan terima kasih kepada dosen pembimbing yang telah memberikan arahan, rekan-rekan yang turut memberikan masukan, serta semua pihak yang tidak dapat disebutkan satu per satu.")
para("Penyusun menyadari sepenuhnya bahwa laporan ini masih memiliki banyak kekurangan. Oleh sebab itu, kritik dan saran yang membangun sangat penyusun harapkan demi penyempurnaan pada masa mendatang. Semoga laporan ini memberikan manfaat bagi pembaca serta bagi pengembangan aplikasi sejenis.")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.add_run("\nPenyusun")

page_break()

# ============================================================
# DAFTAR ISI (placeholder)
# ============================================================
H1("DAFTAR ISI")
para("Daftar isi, daftar tabel, dan daftar gambar disarankan dibuat secara otomatis melalui menu References pada pengolah kata setelah seluruh penomoran halaman ditetapkan. Berikut kerangka pokok laporan sebagai acuan.", indent=False)
for item in [
    "BAB I PENDAHULUAN",
    "BAB II TINJAUAN PUSTAKA",
    "BAB III ANALISIS DAN PERANCANGAN SISTEM",
    "BAB IV IMPLEMENTASI DAN PEMBAHASAN",
    "BAB V PENUTUP",
    "DAFTAR PUSTAKA",
    "LAMPIRAN",
]:
    bullet(item)

page_break()

# ============================================================
# BAB I PENDAHULUAN
# ============================================================
H1("BAB I  PENDAHULUAN")

H2("1.1  Latar Belakang")
para("Toko bangunan merupakan salah satu jenis usaha ritel yang memiliki peran penting dalam menunjang kegiatan pembangunan dan perbaikan infrastruktur di lingkungan masyarakat. Usaha ini menjual beragam jenis barang seperti semen, besi, cat, pipa, paku, keramik, hingga peralatan pertukangan dengan satuan dan harga yang bervariasi. Karakteristik barang yang banyak ragamnya, perputaran stok yang cepat, serta pola pembayaran yang kerap dilakukan secara tunai maupun kredit menjadikan pengelolaan administrasi pada toko bangunan tergolong kompleks apabila tidak ditunjang oleh sistem yang memadai.")
para("Pada praktiknya, sebagian besar toko bangunan berskala kecil dan menengah masih mengandalkan pencatatan transaksi secara manual. Pencatatan dilakukan menggunakan buku tulis, nota rangkap, atau berkas lembar sebar sederhana yang dikelola seadanya. Cara seperti ini memang terlihat praktis pada awalnya, namun seiring meningkatnya volume transaksi, kelemahannya menjadi semakin nyata. Catatan mudah hilang, tulisan sulit dibaca, perhitungan rentan keliru, dan data antara satu buku dengan buku lainnya sering kali tidak selaras.")
para("Kendala lain yang lazim ditemui adalah sulitnya pemilik toko mengetahui kondisi keuangan usahanya secara tepat waktu. Untuk mengetahui laba yang diperoleh dalam suatu periode, pemilik harus menghitung ulang seluruh catatan penjualan, pembelian, dan pengeluaran satu per satu. Proses ini memakan waktu, melelahkan, dan hasilnya belum tentu akurat. Akibatnya, banyak keputusan usaha yang penting, seperti penentuan harga jual, pemilihan barang yang perlu ditambah stoknya, atau penagihan piutang kepada pelanggan, diambil hanya berdasarkan perkiraan dan ingatan, bukan berdasarkan data yang dapat dipertanggungjawabkan.")
para("Perkembangan teknologi informasi membuka peluang untuk mengatasi persoalan tersebut melalui digitalisasi usaha. Digitalisasi pembukuan memungkinkan setiap transaksi dicatat sekali saja, kemudian diolah secara otomatis menjadi informasi stok, laporan keuangan, serta dokumen nota yang rapi. Aplikasi berbasis web menjadi pilihan yang relevan karena dapat diakses melalui peramban pada komputer maupun perangkat bergerak tanpa perlu pemasangan perangkat lunak khusus, sehingga sesuai dengan kebutuhan pelaku usaha yang menginginkan kemudahan dan fleksibilitas.")
para("Berangkat dari kondisi tersebut, dikembangkanlah sebuah aplikasi bernama Bangun Kelola. Aplikasi ini dirancang khusus untuk membantu toko bangunan dalam mengelola transaksi penjualan, pembelian, pencatatan stok, pembukuan, laporan keuangan, serta pencetakan nota secara terpadu. Bangun Kelola berusaha menjawab kebutuhan nyata pelaku usaha dengan menghadirkan istilah dan alur kerja yang akrab dengan dunia toko bangunan, sekaligus memanfaatkan teknologi web modern agar sistem ringan, cepat, dan mudah digunakan oleh pengguna yang tidak memiliki latar belakang teknis.")

H2("1.2  Rumusan Masalah")
para("Berdasarkan uraian latar belakang di atas, rumusan masalah yang diangkat dalam proyek ini adalah sebagai berikut.")
bullet("Bagaimana kondisi pengelolaan pembukuan pada toko bangunan yang masih dilakukan secara manual dan kendala apa saja yang ditimbulkannya?", num=True)
bullet("Bagaimana merancang sebuah aplikasi pembukuan dan akuntansi berbasis web yang sesuai dengan kebutuhan dan alur kerja toko bangunan?", num=True)
bullet("Bagaimana mengimplementasikan fitur pencatatan transaksi penjualan, pembelian, manajemen stok, pembukuan, laporan keuangan, dan pencetakan nota ke dalam aplikasi Bangun Kelola?", num=True)
bullet("Bagaimana aplikasi Bangun Kelola dapat membantu pemilik toko, admin, dan kasir dalam mengelola usaha secara lebih efektif dibandingkan pencatatan manual?", num=True)

H2("1.3  Tujuan")
para("Sejalan dengan rumusan masalah, tujuan yang hendak dicapai melalui pengembangan aplikasi Bangun Kelola adalah sebagai berikut.")
bullet("Mengidentifikasi permasalahan dan kebutuhan pengelolaan pembukuan pada toko bangunan yang masih menggunakan pencatatan manual.", num=True)
bullet("Merancang model aplikasi pembukuan dan akuntansi berbasis web yang menyesuaikan diri dengan proses bisnis toko bangunan.", num=True)
bullet("Membangun dan menerapkan fitur penjualan, pembelian, manajemen stok, pembukuan, laporan keuangan, serta pencetakan nota dalam satu sistem yang terintegrasi.", num=True)
bullet("Menyediakan sarana yang memudahkan pemilik toko, admin, dan kasir dalam mencatat, memantau, dan menganalisis kondisi keuangan usaha secara cepat dan akurat.", num=True)

H2("1.4  Manfaat")
para("Pengembangan aplikasi Bangun Kelola diharapkan memberikan manfaat kepada beberapa pihak sebagaimana diuraikan berikut ini.")
H3("a. Bagi Pemilik Toko")
para("Pemilik toko memperoleh gambaran kondisi usaha yang lebih jelas dan menyeluruh. Melalui ringkasan pada dasbor serta laporan laba rugi, arus kas, dan nilai persediaan, pemilik dapat memantau kinerja usaha kapan saja tanpa harus menghitung ulang catatan secara manual. Informasi mengenai piutang pelanggan, hutang kepada pemasok, dan barang yang stoknya menipis turut membantu pemilik dalam mengambil keputusan yang lebih tepat dan terukur.")
H3("b. Bagi Admin")
para("Admin terbantu dalam mengelola data master seperti data barang, kategori, pemasok, dan pelanggan secara terpusat. Proses pencatatan pembelian, penyesuaian stok, serta penyusunan laporan menjadi lebih ringkas karena seluruh data tersimpan dalam satu basis data yang konsisten. Pekerjaan administrasi yang sebelumnya berulang dan rawan salah dapat dikerjakan dengan lebih cepat dan rapi.")
H3("c. Bagi Kasir")
para("Kasir dapat melayani transaksi penjualan dengan lebih cepat karena harga jual barang sudah tersimpan dan dapat dipanggil secara otomatis. Stok barang langsung berkurang setiap kali terjadi penjualan, sehingga kasir tidak perlu mencatat ulang secara terpisah. Nota penjualan dapat dicetak atau disimpan dalam bentuk dokumen digital sehingga pelayanan kepada pelanggan menjadi lebih profesional.")
H3("d. Bagi Dunia Pendidikan")
para("Bagi dunia pendidikan, laporan dan aplikasi ini dapat menjadi bahan pembelajaran mengenai penerapan sistem informasi akuntansi pada usaha ritel. Proyek ini menunjukkan bagaimana konsep pembukuan dan akuntansi diterjemahkan ke dalam sebuah perangkat lunak nyata dengan memanfaatkan teknologi web modern, sehingga dapat menjadi rujukan maupun pembanding bagi penelitian dan pengembangan sejenis pada masa mendatang.")

H2("1.5  Ruang Lingkup")
para("Agar pembahasan terarah dan tidak meluas, ruang lingkup serta batasan aplikasi Bangun Kelola ditetapkan sebagai berikut.")
bullet("Aplikasi dikembangkan dalam bentuk aplikasi berbasis web yang diakses melalui peramban dan ditujukan secara khusus untuk usaha toko bangunan.")
bullet("Cakupan fitur meliputi autentikasi pengguna, dasbor ringkasan, pengelolaan data barang dan kategori, pengelolaan pemasok dan pelanggan, transaksi pembelian, transaksi penjualan, manajemen stok, penyesuaian dan mutasi stok, pencatatan pengeluaran, pengelolaan piutang dan hutang, pembukuan dan laporan keuangan, ekspor data ke Excel, pencetakan nota dalam format PDF, serta dukungan dua bahasa.")
bullet("Setiap akun pengguna mengelola data tokonya sendiri secara terpisah dan aman, sehingga data milik satu toko tidak dapat diakses oleh toko lain.")
bullet("Laporan keuangan yang disajikan berfokus pada kebutuhan praktis pelaku usaha, yaitu laba rugi, arus kas, posisi keuangan sederhana, umur piutang, dan umur hutang, serta belum mencakup penyusunan laporan keuangan berstandar akuntansi keuangan secara penuh.")
bullet("Pembahasan dalam laporan ini menitikberatkan pada analisis kebutuhan, perancangan, implementasi fitur, manfaat, serta teknologi yang digunakan, dan tidak membahas kode program maupun algoritma pemrograman secara mendalam.")

page_break()

# ============================================================
# BAB II TINJAUAN PUSTAKA
# ============================================================
H1("BAB II  TINJAUAN PUSTAKA")
para("Bab ini menguraikan landasan teori yang menjadi dasar pengembangan aplikasi Bangun Kelola. Teori yang dibahas mencakup konsep pembukuan dan akuntansi, sistem informasi, teknologi web dan basis data, serta perangkat teknologi yang digunakan dalam membangun aplikasi. Pada bagian akhir disajikan pula beberapa penelitian terdahulu yang relevan sebagai pembanding.")

H2("2.1  Pembukuan")
para("Pembukuan adalah kegiatan mencatat secara teratur seluruh transaksi keuangan yang terjadi dalam suatu usaha, meliputi pemasukan, pengeluaran, serta perubahan harta dan kewajiban. Tujuan utama pembukuan adalah menyediakan catatan yang tertib sehingga pelaku usaha dapat mengetahui posisi keuangannya dengan jelas dan dapat mempertanggungjawabkan setiap arus dana yang terjadi.")
para("Dalam konteks usaha kecil seperti toko bangunan, pembukuan yang baik menjadi fondasi pengelolaan usaha yang sehat. Catatan transaksi yang rapi memudahkan pemilik dalam menghitung laba, mengendalikan biaya, serta menelusuri kembali transaksi apabila terjadi perselisihan dengan pelanggan maupun pemasok. Tanpa pembukuan yang tertib, kondisi keuangan usaha menjadi kabur dan sulit dievaluasi.")
para("Pembukuan secara tradisional dilakukan dengan tulisan tangan pada buku kas. Namun seiring perkembangan teknologi, pembukuan kini dapat dilakukan secara digital melalui aplikasi yang mampu mencatat, menghitung, dan menyajikan informasi keuangan secara otomatis. Pergeseran ini membuat pembukuan menjadi lebih cepat, akurat, dan mudah ditelusuri.")

H2("2.2  Akuntansi")
para("Akuntansi merupakan proses mengidentifikasi, mencatat, mengukur, mengelompokkan, mengikhtisarkan, serta melaporkan transaksi ekonomi suatu entitas sehingga menghasilkan informasi keuangan yang berguna bagi pengambilan keputusan. Jika pembukuan lebih menekankan pada kegiatan pencatatan, akuntansi memiliki cakupan yang lebih luas karena turut menafsirkan dan menyajikan informasi keuangan dalam bentuk laporan yang bermakna.")
para("Konsep dasar akuntansi yang relevan dengan aplikasi ini antara lain pendapatan, harga pokok penjualan, beban operasional, laba kotor, dan laba bersih. Pendapatan diperoleh dari penjualan barang, harga pokok penjualan merupakan nilai modal barang yang terjual, sementara beban operasional mencakup pengeluaran untuk menjalankan usaha. Selisih dari unsur-unsur tersebut menghasilkan laba atau rugi pada suatu periode.")
para("Penerapan prinsip akuntansi yang konsisten membuat informasi keuangan dapat dibandingkan antarperiode. Hal ini penting bagi toko bangunan agar pemilik dapat menilai apakah kinerja usahanya membaik atau menurun dari waktu ke waktu, serta dapat mengambil langkah perbaikan secara tepat.")

H2("2.3  Sistem Informasi")
para("Sistem informasi adalah gabungan dari komponen manusia, perangkat keras, perangkat lunak, data, dan prosedur yang bekerja sama untuk mengumpulkan, mengolah, menyimpan, dan menyebarkan informasi guna mendukung kegiatan operasional maupun pengambilan keputusan dalam suatu organisasi. Sistem informasi mengubah data mentah menjadi informasi yang berguna bagi penggunanya.")
para("Pada usaha toko bangunan, sistem informasi berperan sebagai sarana yang menyatukan berbagai data transaksi ke dalam satu wadah yang terorganisasi. Dengan demikian, data penjualan, pembelian, dan stok tidak lagi tersebar pada catatan terpisah, melainkan saling terhubung dan dapat diolah menjadi laporan yang utuh.")
para("Keberadaan sistem informasi yang baik membantu pelaku usaha bekerja lebih efisien. Pekerjaan yang sebelumnya dilakukan secara berulang dan manual dapat diotomasi, sehingga waktu dan tenaga dapat dialihkan untuk kegiatan yang lebih produktif seperti pelayanan pelanggan dan pengembangan usaha.")

H2("2.4  Sistem Informasi Akuntansi")
para("Sistem informasi akuntansi adalah sistem informasi yang secara khusus dirancang untuk mengumpulkan, mencatat, menyimpan, dan mengolah data keuangan sehingga menghasilkan informasi akuntansi bagi pihak yang membutuhkan. Sistem ini menggabungkan prinsip akuntansi dengan teknologi informasi untuk menyajikan laporan keuangan secara cepat dan akurat.")
para("Aplikasi Bangun Kelola pada hakikatnya merupakan penerapan sistem informasi akuntansi pada usaha toko bangunan. Setiap transaksi penjualan, pembelian, dan pengeluaran yang dimasukkan ke dalam aplikasi akan diolah menjadi informasi stok, piutang, hutang, serta laporan laba rugi dan arus kas secara otomatis.")
para("Melalui sistem informasi akuntansi, kesalahan perhitungan dapat ditekan karena seluruh proses dilakukan oleh sistem berdasarkan rumus yang telah ditetapkan. Selain itu, informasi keuangan tersaji secara waktu nyata, sehingga mendukung pengambilan keputusan yang lebih cepat dan berdasar data.")

H2("2.5  Website dan Aplikasi Web")
para("Website atau situs web adalah kumpulan halaman yang saling terhubung dan dapat diakses melalui jaringan internet menggunakan peramban. Pada perkembangannya, website tidak lagi sekadar menyajikan informasi statis, melainkan telah berevolusi menjadi aplikasi web yang interaktif dan mampu menjalankan fungsi-fungsi kompleks layaknya perangkat lunak konvensional.")
para("Aplikasi web memiliki keunggulan berupa kemudahan akses karena dapat dibuka dari berbagai perangkat tanpa perlu pemasangan khusus. Pengguna cukup membuka peramban dan masuk ke alamat aplikasi untuk mulai bekerja. Hal ini sangat menguntungkan bagi pelaku usaha yang menginginkan fleksibilitas dalam mengelola usahanya.")
para("Bangun Kelola dibangun sebagai aplikasi web modern yang mengutamakan kecepatan dan kenyamanan penggunaan. Antarmuka yang responsif memungkinkan aplikasi tampil dengan baik pada layar komputer maupun perangkat bergerak, sehingga dapat digunakan sesuai kebutuhan di meja kasir maupun saat memeriksa stok di gudang.")

H2("2.6  Basis Data")
para("Basis data adalah kumpulan data yang terorganisasi dan saling berkaitan, yang disimpan secara sistematis sehingga dapat diakses, dikelola, dan diperbarui dengan mudah. Basis data menjadi inti dari hampir setiap sistem informasi karena berfungsi sebagai tempat penyimpanan seluruh data yang diolah oleh aplikasi.")
para("Sistem manajemen basis data bertugas mengatur penyimpanan dan pengambilan data agar tetap konsisten, aman, dan terhindar dari duplikasi yang tidak perlu. Melalui penataan data dalam bentuk tabel-tabel yang saling berelasi, basis data memastikan bahwa setiap informasi tersimpan secara terstruktur dan dapat dipanggil kembali ketika dibutuhkan.")
para("Pada aplikasi Bangun Kelola, basis data menyimpan seluruh data penting seperti data barang, pelanggan, pemasok, transaksi penjualan, pembelian, dan pengeluaran. Relasi antartabel yang dirancang dengan baik memungkinkan aplikasi menyajikan informasi yang saling terhubung, misalnya menghubungkan suatu penjualan dengan barang dan pelanggannya.")

H2("2.7  PostgreSQL")
para("PostgreSQL adalah sistem manajemen basis data relasional bersifat sumber terbuka yang dikenal andal, stabil, dan kaya fitur. PostgreSQL mendukung standar bahasa kueri terstruktur secara luas serta menyediakan kemampuan lanjutan seperti integritas data, transaksi yang aman, dan fungsi yang dapat ditanam langsung pada basis data.")
para("Dalam aplikasi Bangun Kelola, PostgreSQL berperan sebagai basis data utama yang menyimpan seluruh data transaksi dan data master. Kemampuannya dalam menjaga integritas data sangat penting untuk memastikan bahwa angka stok, total transaksi, dan saldo keuangan selalu konsisten meskipun terjadi banyak transaksi dalam waktu bersamaan.")
para("PostgreSQL juga mendukung penggunaan aturan dan pemicu yang dapat menjalankan logika tertentu secara otomatis di tingkat basis data. Pada aplikasi ini, kemampuan tersebut dimanfaatkan untuk memperbarui stok barang secara otomatis ketika terjadi penjualan atau pembelian, serta untuk menjaga agar stok tidak menjadi negatif.")

H2("2.8  Supabase")
para("Supabase adalah platform pengembangan aplikasi sumber terbuka yang menyediakan layanan backend secara siap pakai dengan PostgreSQL sebagai basis datanya. Supabase menawarkan berbagai layanan seperti autentikasi pengguna, antarmuka pemrograman aplikasi yang dihasilkan secara otomatis dari struktur basis data, penyimpanan berkas, serta kemampuan data waktu nyata.")
para("Dengan memanfaatkan Supabase, pengembang tidak perlu membangun seluruh komponen backend dari nol. Layanan autentikasi yang tersedia memudahkan pengelolaan proses masuk dan pendaftaran pengguna, sementara mekanisme keamanan tingkat baris memastikan setiap pengguna hanya dapat mengakses data miliknya sendiri.")
para("Pada aplikasi Bangun Kelola, Supabase berperan sebagai tulang punggung backend yang menangani penyimpanan data, autentikasi, dan pengamanan akses. Pendekatan ini mempercepat proses pengembangan sekaligus menjaga keamanan data setiap toko yang menggunakan aplikasi.")

H2("2.9  React")
para("React adalah pustaka JavaScript yang digunakan untuk membangun antarmuka pengguna, terutama untuk aplikasi web satu halaman. React mengusung konsep komponen, yaitu bagian antarmuka yang dapat digunakan kembali dan dikelola secara mandiri, sehingga memudahkan pengembangan aplikasi yang besar dan kompleks.")
para("Keunggulan React terletak pada kemampuannya memperbarui tampilan secara efisien ketika data berubah. Ketika pengguna melakukan transaksi, React hanya memperbarui bagian antarmuka yang relevan tanpa memuat ulang seluruh halaman, sehingga aplikasi terasa cepat dan responsif.")
para("Aplikasi Bangun Kelola dibangun menggunakan React versi terbaru yang dipadukan dengan kerangka kerja pendukung seperti TanStack Start dan TanStack Router untuk pengaturan rute serta React Query untuk pengelolaan data dari server. Kombinasi ini menghasilkan aplikasi web yang modern, terstruktur, dan nyaman digunakan.")

H2("2.10  TypeScript")
para("TypeScript adalah bahasa pemrograman yang dikembangkan di atas JavaScript dengan penambahan sistem tipe statis. Melalui penambahan tipe, TypeScript membantu pengembang mendeteksi kesalahan sejak tahap penulisan kode, sebelum aplikasi dijalankan, sehingga kualitas dan keandalan perangkat lunak meningkat.")
para("Penggunaan TypeScript sangat bermanfaat pada aplikasi yang menangani data keuangan seperti Bangun Kelola. Dengan tipe data yang jelas, risiko kesalahan dalam pengolahan angka dan struktur data dapat ditekan, sehingga perhitungan total transaksi maupun laporan keuangan menjadi lebih dapat dipercaya.")
para("Selain meningkatkan keandalan, TypeScript juga membuat kode lebih mudah dipahami dan dipelihara. Struktur data yang terdefinisi dengan baik memudahkan pengembangan fitur baru maupun perbaikan pada masa mendatang tanpa mengganggu bagian lain dari aplikasi.")

H2("2.11  Tailwind CSS")
para("Tailwind CSS adalah kerangka kerja penataan gaya yang menerapkan pendekatan berbasis kelas utilitas. Alih-alih menulis berkas gaya yang terpisah dan panjang, pengembang menyusun tampilan langsung melalui kelas-kelas ringkas yang masing-masing mengatur satu aspek tampilan, seperti warna, jarak, atau ukuran.")
para("Pendekatan ini mempercepat proses pembuatan antarmuka serta menjaga konsistensi tampilan di seluruh halaman aplikasi. Tailwind CSS juga memudahkan pembuatan desain yang responsif, sehingga tampilan aplikasi dapat menyesuaikan diri secara otomatis terhadap berbagai ukuran layar.")
para("Pada aplikasi Bangun Kelola, Tailwind CSS dipadukan dengan komponen antarmuka dari Radix UI serta ikon dari Lucide React untuk menghasilkan tampilan yang bersih, modern, dan mudah digunakan. Grafik visual pada dasbor dan laporan dibangun menggunakan pustaka Recharts agar data dapat disajikan dalam bentuk yang mudah dipahami.")

H2("2.12  Konsep Pembukuan Digital")
para("Pembukuan digital merupakan praktik pencatatan keuangan yang dilakukan dengan bantuan perangkat lunak atau aplikasi, menggantikan pencatatan manual berbasis kertas. Konsep ini menekankan pada otomatisasi, integrasi data, dan penyajian informasi yang cepat sebagai pengganti proses pencatatan dan perhitungan yang sebelumnya dilakukan secara manual.")
para("Keunggulan utama pembukuan digital adalah prinsip satu kali masuk untuk banyak keluaran. Sebuah transaksi yang dicatat sekali dapat langsung memengaruhi stok, kas, piutang, dan laporan keuangan secara bersamaan. Hal ini menghilangkan kebutuhan untuk menulis ulang data yang sama di beberapa tempat, sekaligus mengurangi risiko ketidaksesuaian data.")
para("Bangun Kelola menerapkan konsep pembukuan digital secara menyeluruh. Setiap transaksi yang dimasukkan akan diproses menjadi berbagai informasi yang saling terkait, sehingga pemilik toko memperoleh gambaran usaha yang utuh tanpa harus melakukan perhitungan tambahan secara manual.")

H2("2.13  Laporan Keuangan")
para("Laporan keuangan adalah hasil akhir dari proses akuntansi yang menyajikan informasi mengenai posisi keuangan dan kinerja suatu usaha dalam suatu periode tertentu. Laporan keuangan menjadi alat penting bagi pemilik usaha untuk menilai keberhasilan usaha serta sebagai dasar pengambilan keputusan.")
para("Beberapa bentuk laporan keuangan yang umum antara lain laporan laba rugi yang menggambarkan pendapatan dan beban, laporan arus kas yang menggambarkan aliran masuk dan keluar uang, serta laporan posisi keuangan yang menggambarkan harta dan kewajiban. Ketiga laporan tersebut saling melengkapi dalam memberikan gambaran kondisi usaha.")
para("Aplikasi Bangun Kelola menyajikan laporan keuangan yang disesuaikan dengan kebutuhan praktis toko bangunan, meliputi laporan laba rugi, arus kas, posisi keuangan sederhana, laba per kategori barang, serta analisis umur piutang dan umur hutang. Penyajian dalam bentuk angka maupun grafik membantu pengguna memahami kondisi keuangan dengan lebih mudah.")

H2("2.14  Nota Digital")
para("Nota digital adalah dokumen bukti transaksi yang dibuat dan disimpan dalam bentuk berkas elektronik, umumnya dalam format PDF. Berbeda dengan nota kertas yang mudah rusak atau hilang, nota digital dapat disimpan, dicetak ulang, dan dibagikan dengan mudah kapan pun diperlukan.")
para("Penggunaan nota digital memberikan kesan profesional sekaligus memudahkan penelusuran transaksi. Setiap nota memuat informasi lengkap seperti nomor nota, tanggal, daftar barang, jumlah, harga, dan total pembayaran, sehingga dapat menjadi rujukan yang jelas baik bagi penjual maupun pembeli.")
para("Bangun Kelola menyediakan fitur pembuatan nota digital dalam format PDF melalui pustaka jsPDF. Nota dapat dicetak langsung atau disimpan sebagai berkas, sehingga toko dapat memberikan bukti transaksi yang rapi kepada pelanggan tanpa bergantung sepenuhnya pada nota kertas.")

H2("2.15  Penelitian Terdahulu")
para("Beberapa penelitian terdahulu yang relevan dengan pengembangan sistem informasi pembukuan dan akuntansi untuk usaha ritel dirangkum pada tabel berikut sebagai bahan pembanding. Penelitian-penelitian tersebut menunjukkan bahwa penerapan sistem informasi dapat meningkatkan efektivitas pengelolaan usaha dibandingkan pencatatan manual.")
table(
    ["No", "Topik Penelitian", "Fokus", "Relevansi dengan Bangun Kelola"],
    [
        ["1", "Sistem informasi penjualan pada usaha ritel berbasis web", "Pencatatan penjualan dan stok secara terkomputerisasi", "Menjadi acuan pada fitur penjualan dan manajemen stok"],
        ["2", "Aplikasi pembukuan keuangan untuk usaha mikro, kecil, dan menengah", "Penyusunan laporan keuangan sederhana secara otomatis", "Mendasari fitur pembukuan dan laporan laba rugi"],
        ["3", "Sistem informasi akuntansi persediaan barang dagang", "Pengendalian persediaan dan harga pokok", "Menjadi rujukan perhitungan stok dan harga pokok penjualan"],
        ["4", "Penerapan nota elektronik pada transaksi penjualan", "Pembuatan bukti transaksi digital", "Mendasari fitur pencetakan nota dalam format PDF"],
    ],
    widths=[0.4, 2.1, 2.0, 2.2],
)
para("Berbeda dengan penelitian terdahulu yang umumnya bersifat umum, aplikasi Bangun Kelola menonjolkan kekhususan pada domain toko bangunan. Aplikasi ini memadukan fitur penjualan, pembelian, manajemen stok, pembukuan, laporan keuangan, dan nota digital dalam satu sistem dengan istilah serta alur kerja yang akrab bagi pelaku usaha bahan bangunan.")

page_break()

# ============================================================
# BAB III ANALISIS DAN PERANCANGAN SISTEM
# ============================================================
H1("BAB III  ANALISIS DAN PERANCANGAN SISTEM")

H2("3.1  Gambaran Umum Bangun Kelola")
para("Bangun Kelola adalah aplikasi pembukuan dan akuntansi berbasis web yang dirancang khusus untuk membantu pengelolaan usaha toko bangunan. Aplikasi ini menyatukan seluruh kegiatan administrasi toko, mulai dari pencatatan barang masuk melalui pembelian, pencatatan barang keluar melalui penjualan, pengelolaan stok, hingga penyusunan laporan keuangan, ke dalam satu sistem yang terintegrasi dan mudah digunakan.")
para("Aplikasi ini bekerja dengan prinsip bahwa setiap transaksi cukup dimasukkan satu kali, kemudian sistem secara otomatis memperbarui stok barang, menghitung total transaksi, mencatat piutang atau hutang, serta memperbarui laporan keuangan. Dengan demikian, pengguna tidak perlu melakukan perhitungan manual yang berulang dan rawan kesalahan.")
para("Setiap pengguna yang mendaftar akan memiliki ruang data tersendiri yang terpisah dari pengguna lain. Hal ini memastikan bahwa data setiap toko tersimpan secara aman dan tidak tercampur. Aplikasi juga menyediakan dukungan dua bahasa, yaitu Bahasa Indonesia dan Bahasa Inggris, sehingga dapat digunakan oleh kalangan pengguna yang lebih luas.")
para("Secara umum, Bangun Kelola dirancang agar dapat dioperasikan oleh pengguna yang tidak memiliki latar belakang teknis. Istilah yang digunakan disesuaikan dengan kebiasaan pelaku usaha toko bangunan, dan alur kerja aplikasi dibuat sederhana agar mudah dipelajari dalam waktu singkat.")

H2("3.2  Analisis Kebutuhan Pengguna")
para("Analisis kebutuhan pengguna dilakukan untuk memahami pihak-pihak yang akan menggunakan aplikasi serta kebutuhan masing-masing. Pada usaha toko bangunan, terdapat tiga peran utama yang berinteraksi dengan sistem, yaitu pemilik toko, admin, dan kasir. Ketiga peran tersebut memiliki kebutuhan yang berbeda namun saling melengkapi.")
para("Pemilik toko membutuhkan informasi menyeluruh mengenai kondisi usahanya. Kebutuhan utamanya adalah memantau laba rugi, arus kas, nilai persediaan, serta posisi piutang dan hutang. Pemilik juga memerlukan kemampuan untuk meninjau kinerja barang dan mengevaluasi usaha dari waktu ke waktu sebagai dasar pengambilan keputusan.")
para("Admin membutuhkan kemudahan dalam mengelola data master dan transaksi. Kebutuhan admin meliputi pencatatan data barang, kategori, pemasok, dan pelanggan, pencatatan pembelian, penyesuaian stok, serta penyusunan dan ekspor laporan. Sementara itu, kasir membutuhkan proses pencatatan penjualan yang cepat dan akurat, termasuk kemampuan mencetak nota bagi pelanggan.")
para("Berdasarkan analisis tersebut, kebutuhan masing-masing pengguna dirangkum dalam tabel berikut.")
table(
    ["Pengguna", "Kebutuhan Utama"],
    [
        ["Pemilik Toko", "Memantau dasbor, laporan laba rugi, arus kas, nilai persediaan, piutang, hutang, dan kinerja barang."],
        ["Admin", "Mengelola data barang, kategori, pemasok, pelanggan, transaksi pembelian, penyesuaian stok, serta menyusun dan mengekspor laporan."],
        ["Kasir", "Mencatat transaksi penjualan secara cepat, memeriksa stok, dan mencetak nota penjualan."],
    ],
    widths=[1.6, 5.1],
)

H2("3.3  Analisis Kebutuhan Sistem")
para("Analisis kebutuhan sistem dibagi menjadi dua, yaitu kebutuhan fungsional dan kebutuhan nonfungsional. Kebutuhan fungsional menggambarkan fungsi yang harus dapat dijalankan oleh aplikasi, sedangkan kebutuhan nonfungsional menggambarkan kualitas dan karakteristik yang harus dipenuhi sistem.")
H3("a. Kebutuhan Fungsional")
table(
    ["Kode", "Kebutuhan Fungsional"],
    [
        ["KF-01", "Sistem dapat melakukan pendaftaran dan autentikasi pengguna."],
        ["KF-02", "Sistem dapat menampilkan dasbor ringkasan kondisi usaha."],
        ["KF-03", "Sistem dapat mengelola data barang dan kategori."],
        ["KF-04", "Sistem dapat menetapkan stok awal barang."],
        ["KF-05", "Sistem dapat mengelola data pemasok dan pelanggan."],
        ["KF-06", "Sistem dapat mencatat transaksi pembelian dan menambah stok secara otomatis."],
        ["KF-07", "Sistem dapat mencatat transaksi penjualan dan mengurangi stok secara otomatis."],
        ["KF-08", "Sistem dapat melakukan penyesuaian dan menampilkan mutasi stok."],
        ["KF-09", "Sistem dapat mencatat pengeluaran usaha."],
        ["KF-10", "Sistem dapat mengelola piutang pelanggan dan hutang kepada pemasok."],
        ["KF-11", "Sistem dapat menyajikan laporan keuangan seperti laba rugi dan arus kas."],
        ["KF-12", "Sistem dapat mengekspor data ke berkas Excel dan mencetak nota dalam format PDF."],
        ["KF-13", "Sistem dapat menampilkan antarmuka dalam dua bahasa."],
    ],
    widths=[1.0, 5.7],
)
H3("b. Kebutuhan Nonfungsional")
table(
    ["Aspek", "Kebutuhan Nonfungsional"],
    [
        ["Keamanan", "Setiap pengguna hanya dapat mengakses data miliknya sendiri melalui mekanisme keamanan tingkat baris."],
        ["Kemudahan", "Antarmuka mudah dipahami dan dapat digunakan tanpa pelatihan teknis khusus."],
        ["Ketersediaan", "Aplikasi dapat diakses melalui peramban kapan saja selama terhubung internet."],
        ["Kinerja", "Aplikasi memperbarui tampilan dengan cepat tanpa memuat ulang seluruh halaman."],
        ["Keterpakaian", "Tampilan responsif dan dapat digunakan pada komputer maupun perangkat bergerak."],
    ],
    widths=[1.4, 5.3],
)

H2("3.4  Analisis Proses Bisnis Toko Bangunan")
para("Proses bisnis toko bangunan secara umum berputar di sekitar arus barang dan arus uang. Arus barang dimulai dari pembelian barang kepada pemasok, kemudian barang disimpan sebagai stok, dan akhirnya dijual kepada pelanggan. Arus uang mengikuti arus barang tersebut, baik berupa pembayaran kepada pemasok maupun penerimaan dari pelanggan.")
para("Pada sisi pembelian, toko memesan barang kepada pemasok, menerima barang, dan mencatat penambahan stok. Pembayaran kepada pemasok dapat dilakukan secara tunai maupun kredit. Apabila dilakukan secara kredit, timbul hutang yang harus dilunasi sesuai tanggal jatuh tempo yang disepakati.")
para("Pada sisi penjualan, pelanggan membeli barang dan toko mencatat pengurangan stok. Pembayaran dari pelanggan juga dapat dilakukan secara tunai maupun kredit. Penjualan kredit menimbulkan piutang yang perlu ditagih sesuai tenggat yang berlaku, sehingga toko perlu memantau umur piutang agar tidak menumpuk.")
para("Selain transaksi pembelian dan penjualan, toko juga menanggung berbagai pengeluaran operasional seperti gaji, sewa, transportasi, dan biaya lainnya. Seluruh komponen tersebut, yaitu pendapatan penjualan, harga pokok barang, dan beban operasional, pada akhirnya menentukan laba atau rugi usaha. Bangun Kelola memetakan keseluruhan proses bisnis ini ke dalam fitur-fiturnya secara terstruktur.")

H2("3.5  Flow Proses Aplikasi")
para("Alur penggunaan aplikasi Bangun Kelola dirancang mengikuti urutan kerja yang wajar pada toko bangunan. Pengguna memulai dengan masuk ke aplikasi, kemudian menyiapkan data dasar, melakukan transaksi harian, dan diakhiri dengan peninjauan laporan. Gambaran alur tersebut diuraikan sebagai berikut.")
bullet("Pengguna mendaftar atau masuk ke aplikasi menggunakan email dan kata sandi.", num=True)
bullet("Pengguna menyiapkan data master, yaitu data barang beserta stok awalnya, kategori, pemasok, dan pelanggan.", num=True)
bullet("Pengguna mencatat transaksi pembelian sehingga stok barang bertambah, atau mencatat transaksi penjualan sehingga stok berkurang.", num=True)
bullet("Pengguna mencatat pengeluaran operasional serta melakukan penyesuaian stok bila diperlukan.", num=True)
bullet("Pengguna memantau piutang dan hutang, lalu mencatat pembayaran ketika terjadi pelunasan.", num=True)
bullet("Pengguna meninjau dasbor dan laporan keuangan, kemudian mengekspor data atau mencetak nota sesuai kebutuhan.", num=True)
figure_placeholder("Gambar 3.1 Diagram Alur (Flowchart) Proses Penggunaan Aplikasi Bangun Kelola")
caption("Gambar 3.1 Diagram alur proses penggunaan aplikasi")

H2("3.6  Use Case Diagram")
para("Use case diagram menggambarkan interaksi antara pengguna dengan fungsi-fungsi yang disediakan oleh aplikasi. Pada Bangun Kelola, ketiga peran pengguna dapat dilayani oleh satu akun pemilik usaha, namun secara konseptual kebutuhan tiap peran dapat dipetakan agar lebih jelas. Tabel berikut menyajikan pemetaan aktor dan use case utama.")
table(
    ["Aktor", "Use Case yang Dapat Diakses"],
    [
        ["Pemilik Toko", "Melihat dasbor, melihat laporan keuangan, meninjau piutang dan hutang, mengatur pengaturan, mengelola seluruh data."],
        ["Admin", "Mengelola data barang dan kategori, mengelola pemasok dan pelanggan, mencatat pembelian, menyesuaikan stok, menyusun laporan."],
        ["Kasir", "Mencatat penjualan, memeriksa stok, mencetak nota."],
    ],
    widths=[1.6, 5.1],
)
figure_placeholder("Gambar 3.2 Use Case Diagram Aplikasi Bangun Kelola")
caption("Gambar 3.2 Use case diagram aplikasi")
para("Diagram tersebut menunjukkan bahwa seluruh fungsi aplikasi pada dasarnya berpusat pada pengelolaan transaksi dan penyajian informasi. Aktor pemilik memiliki cakupan akses paling luas, sedangkan admin dan kasir berfokus pada kegiatan operasional sehari-hari sesuai perannya masing-masing.")

H2("3.7  Activity Diagram")
para("Activity diagram menggambarkan urutan aktivitas dalam menjalankan suatu proses pada aplikasi. Sebagai contoh, proses pencatatan transaksi penjualan dapat digambarkan sebagai rangkaian aktivitas yang berurutan dari awal hingga selesai.")
para("Proses dimulai ketika kasir membuka halaman penjualan dan menekan tombol untuk menambah transaksi baru. Kasir kemudian memilih pelanggan apabila diperlukan, menambahkan barang beserta jumlahnya, dan sistem menampilkan total secara otomatis. Setelah itu kasir memilih metode pembayaran dan memasukkan jumlah yang dibayarkan.")
para("Sebelum transaksi disimpan, sistem memeriksa ketersediaan stok dan, untuk penjualan kredit, memeriksa batas kredit pelanggan. Apabila pemeriksaan berhasil, transaksi disimpan, stok barang berkurang secara otomatis, dan apabila pembayaran belum lunas maka sisa pembayaran dicatat sebagai piutang. Proses diakhiri dengan tersedianya nota yang dapat dicetak.")
figure_placeholder("Gambar 3.3 Activity Diagram Proses Transaksi Penjualan")
caption("Gambar 3.3 Activity diagram proses transaksi penjualan")

H2("3.8  Entity Relationship Diagram (ERD)")
para("Entity Relationship Diagram menggambarkan entitas data yang terdapat dalam sistem beserta hubungan antarentitas. Pada aplikasi Bangun Kelola, entitas utama meliputi pengguna, kategori, barang, pelanggan, pemasok, penjualan beserta rinciannya, pembelian beserta rinciannya, pengeluaran, dan penyesuaian stok.")
para("Hubungan antarentitas dibangun mengikuti logika proses bisnis. Satu barang dapat dimiliki oleh satu kategori, satu penjualan dapat memuat banyak rincian barang, dan satu pembelian juga dapat memuat banyak rincian barang. Pelanggan terhubung dengan transaksi penjualan, sedangkan pemasok terhubung dengan transaksi pembelian.")
para("Seluruh entitas terikat pada entitas pengguna sebagai pemilik data. Keterikatan ini memastikan bahwa setiap data hanya menjadi milik dan dapat diakses oleh pengguna yang bersangkutan. Struktur relasi yang demikian menjaga keterhubungan dan integritas data di seluruh aplikasi.")
figure_placeholder("Gambar 3.4 Entity Relationship Diagram Aplikasi Bangun Kelola")
caption("Gambar 3.4 Entity Relationship Diagram aplikasi")

H2("3.9  Struktur Tabel Basis Data")
para("Basis data aplikasi Bangun Kelola terdiri atas sejumlah tabel yang saling berelasi. Tabel-tabel tersebut menyimpan data master maupun data transaksi. Rincian tabel utama beserta fungsinya disajikan pada tabel berikut, diikuti dengan contoh struktur beberapa tabel inti.")
table(
    ["Nama Tabel", "Fungsi"],
    [
        ["profiles", "Menyimpan data profil pengguna dan nama toko."],
        ["categories", "Menyimpan data kategori barang."],
        ["products", "Menyimpan data barang beserta stok, harga pokok, dan harga jual."],
        ["customers", "Menyimpan data pelanggan beserta batas dan tempo kredit."],
        ["suppliers", "Menyimpan data pemasok."],
        ["sales", "Menyimpan data transaksi penjualan."],
        ["sale_items", "Menyimpan rincian barang pada setiap penjualan."],
        ["purchases", "Menyimpan data transaksi pembelian."],
        ["purchase_items", "Menyimpan rincian barang pada setiap pembelian."],
        ["expenses", "Menyimpan data pengeluaran operasional."],
        ["stock_adjustments", "Menyimpan data penyesuaian stok barang."],
    ],
    widths=[1.8, 4.9],
)
para("Sebagai gambaran rinci, berikut adalah struktur tabel barang (products) yang menjadi salah satu tabel inti dalam aplikasi.")
table(
    ["Kolom", "Tipe", "Keterangan"],
    [
        ["id", "UUID", "Pengenal unik barang."],
        ["user_id", "UUID", "Pengenal pemilik data."],
        ["category_id", "UUID", "Relasi ke kategori barang."],
        ["name", "Teks", "Nama barang."],
        ["sku", "Teks", "Kode barang."],
        ["unit", "Teks", "Satuan barang."],
        ["stock", "Numerik", "Jumlah stok saat ini."],
        ["min_stock", "Numerik", "Batas stok minimum."],
        ["cost_price", "Numerik", "Harga pokok per satuan."],
        ["sell_price", "Numerik", "Harga jual per satuan."],
        ["rack_location", "Teks", "Lokasi rak penyimpanan."],
        ["opening_stock", "Numerik", "Stok awal barang."],
        ["opening_locked", "Boolean", "Penanda penguncian stok awal."],
    ],
    widths=[1.6, 1.2, 3.9],
)
para("Struktur tabel penjualan (sales) yang menjadi tabel transaksi utama disajikan pula sebagai berikut.")
table(
    ["Kolom", "Tipe", "Keterangan"],
    [
        ["id", "UUID", "Pengenal unik penjualan."],
        ["user_id", "UUID", "Pengenal pemilik data."],
        ["customer_id", "UUID", "Relasi ke pelanggan."],
        ["invoice_no", "Teks", "Nomor nota penjualan."],
        ["sale_date", "Tanggal", "Tanggal penjualan."],
        ["subtotal", "Numerik", "Jumlah sebelum potongan."],
        ["discount", "Numerik", "Potongan harga."],
        ["delivery_fee", "Numerik", "Biaya pengiriman."],
        ["total", "Numerik", "Total akhir transaksi."],
        ["paid", "Numerik", "Jumlah yang telah dibayar."],
        ["payment_method", "Teks", "Metode pembayaran."],
        ["due_date", "Tanggal", "Tanggal jatuh tempo untuk kredit."],
    ],
    widths=[1.6, 1.2, 3.9],
)

H2("3.10  Desain Antarmuka")
para("Desain antarmuka aplikasi Bangun Kelola mengutamakan kesederhanaan, kerapian, dan kemudahan penggunaan. Tata letak dirancang dengan menu navigasi di sisi kiri pada tampilan komputer dan menu yang dapat dilipat pada tampilan perangkat bergerak. Susunan ini memudahkan pengguna berpindah antarhalaman dengan cepat.")
para("Menu navigasi dikelompokkan menurut alur kerja, yaitu kelompok utama yang memuat dasbor, kelompok hulu yang memuat pembelian, barang, penyesuaian, dan mutasi stok, kelompok hilir yang memuat penjualan, piutang, kontak, dan pengeluaran, kelompok finansial yang memuat laporan, serta kelompok akun yang memuat pengaturan dan panduan. Pengelompokan ini membantu pengguna memahami posisi setiap fitur dalam alur usaha.")
para("Warna dan elemen visual dipilih agar nyaman dipandang dan tidak membingungkan. Informasi penting seperti stok yang menipis atau status piutang yang melewati tempo ditandai dengan warna khusus agar mudah dikenali. Grafik pada dasbor dan laporan disajikan dengan bentuk yang sederhana sehingga data dapat dipahami secara sekilas.")
para("Seluruh halaman dirancang responsif sehingga dapat menyesuaikan diri dengan ukuran layar perangkat. Formulir pengisian data dibuat ringkas dengan penanda pada kolom yang wajib diisi, serta dilengkapi pemberitahuan singkat ketika terjadi keberhasilan maupun kesalahan dalam menyimpan data.")
figure_placeholder("Gambar 3.5 Rancangan Tata Letak Antarmuka Aplikasi")
caption("Gambar 3.5 Rancangan tata letak antarmuka aplikasi")

H2("3.11  Penjelasan Setiap Halaman")
para("Aplikasi Bangun Kelola terdiri atas sejumlah halaman yang masing-masing memiliki fungsi tertentu. Penjelasan ringkas setiap halaman disajikan pada tabel berikut, sementara uraian implementasinya dibahas lebih lanjut pada bab berikutnya.")
table(
    ["Halaman", "Fungsi"],
    [
        ["Halaman Awal", "Memperkenalkan aplikasi dan menyediakan pintu masuk untuk mendaftar atau masuk."],
        ["Masuk / Daftar", "Tempat pengguna melakukan autentikasi atau pendaftaran akun baru."],
        ["Dasbor", "Menampilkan ringkasan kondisi usaha berupa saldo kas, piutang, hutang, nilai persediaan, dan grafik."],
        ["Data Barang", "Mengelola daftar barang, kategori, stok awal, harga pokok, dan harga jual."],
        ["Kontak", "Mengelola data pelanggan dan pemasok dalam tampilan bertab."],
        ["Pembelian", "Mencatat transaksi pembelian barang kepada pemasok."],
        ["Penjualan", "Mencatat transaksi penjualan barang kepada pelanggan."],
        ["Penyesuaian Stok", "Mencatat koreksi stok akibat kerusakan, kehilangan, atau hasil opname."],
        ["Mutasi Stok", "Menampilkan pergerakan stok masuk, keluar, dan penyesuaian dalam suatu periode."],
        ["Pengeluaran", "Mencatat biaya operasional usaha."],
        ["Piutang & Hutang", "Mengelola dan mencatat pembayaran piutang pelanggan serta hutang pemasok."],
        ["Laporan", "Menyajikan laporan laba rugi, arus kas, posisi keuangan, dan analisis umur piutang dan hutang."],
        ["Pengaturan", "Mengatur profil, bahasa, dan pengelolaan data."],
        ["Panduan", "Menyediakan petunjuk penggunaan aplikasi."],
    ],
    widths=[1.7, 5.0],
)

page_break()

# ============================================================
# BAB IV IMPLEMENTASI DAN PEMBAHASAN
# ============================================================
H1("BAB IV  IMPLEMENTASI DAN PEMBAHASAN")
para("Bab ini menguraikan hasil penerapan rancangan ke dalam aplikasi yang nyata. Pembahasan mencakup implementasi setiap fitur, cara kerja aplikasi secara umum, peran dan alasan pemilihan teknologi, hasil implementasi, pengujian aplikasi, serta analisis kelebihan dan kekurangan. Sesuai dengan ruang lingkup, pembahasan tidak menguraikan kode program melainkan berfokus pada fungsi dan manfaat tiap fitur.")

H2("4.1  Lingkungan Implementasi")
para("Aplikasi Bangun Kelola dibangun menggunakan bahasa pemrograman TypeScript dengan pustaka antarmuka React versi 19 yang dipadukan dengan kerangka kerja TanStack Start dan TanStack Router. Pengelolaan data dari server ditangani oleh React Query, sementara tampilan ditata menggunakan Tailwind CSS bersama komponen Radix UI, ikon Lucide React, dan grafik Recharts.")
para("Pada sisi backend, aplikasi memanfaatkan Supabase yang menyediakan layanan autentikasi dan basis data PostgreSQL. Formulir dikelola menggunakan React Hook Form dengan validasi data melalui Zod. Fitur tambahan dilengkapi pustaka jsPDF untuk pembuatan nota PDF, XLSX untuk ekspor Excel, i18next untuk dukungan banyak bahasa, dan date-fns untuk pengolahan tanggal.")
para("Pengembangan dikelola menggunakan pengelola paket Bun, sedangkan penerapan aplikasi dilakukan melalui platform Vercel sehingga aplikasi dapat diakses secara daring melalui peramban. Kombinasi perangkat ini dipilih untuk menghasilkan aplikasi web yang modern, ringan, dan mudah dipelihara.")

H2("4.2  Implementasi Fitur Aplikasi")

H3("4.2.1  Login dan Pendaftaran")
para("Fitur login dan pendaftaran menjadi pintu masuk aplikasi. Pada halaman ini, pengguna baru dapat mendaftar dengan mengisi nama toko, nama lengkap, email, dan kata sandi, sedangkan pengguna lama cukup memasukkan email dan kata sandi untuk masuk. Setelah berhasil, pengguna langsung diarahkan ke halaman utama aplikasi.")
para("Proses autentikasi ditangani oleh layanan Supabase yang menjaga keamanan akun. Begitu pengguna mendaftar, sistem secara otomatis membuatkan profil toko sehingga pengguna dapat langsung menggunakan aplikasi tanpa pengaturan tambahan yang rumit.")

H3("4.2.2  Dashboard")
para("Dasbor menyajikan ringkasan kondisi usaha secara waktu nyata. Pada bagian atas ditampilkan kartu informasi berupa saldo kas bersih, piutang usaha, hutang usaha, dan nilai persediaan dalam rentang tiga puluh hari terakhir. Informasi ini membantu pengguna memperoleh gambaran cepat mengenai kesehatan usahanya.")
para("Di bawah kartu informasi, dasbor menampilkan grafik tren penjualan dan pengeluaran, daftar barang yang stoknya menipis, grafik barang terlaris, serta komposisi penjualan dalam bentuk diagram lingkaran. Penyajian visual ini memudahkan pengguna memahami pola usaha tanpa harus membaca angka satu per satu.")

H3("4.2.3  Data Barang")
para("Halaman data barang digunakan untuk mengelola seluruh barang dagangan. Pengguna dapat menambah, mengubah, dan menghapus barang, serta mengatur kategori, satuan, lokasi rak, harga pokok, dan harga jual. Daftar barang dilengkapi fitur pencarian berdasarkan nama atau kode barang.")
para("Setiap barang baru dapat ditetapkan stok awalnya melalui fitur stok awal. Setelah barang memiliki transaksi, stok awal akan terkunci agar data historis tetap terjaga. Barang dengan stok di bawah batas minimum ditandai secara khusus agar pengguna segera mengetahui barang yang perlu ditambah.")

H3("4.2.4  Supplier dan Pelanggan")
para("Pengelolaan pemasok dan pelanggan disatukan dalam halaman kontak yang menggunakan tampilan bertab. Pada tab pelanggan, pengguna dapat menyimpan nama, nomor telepon, alamat, batas kredit, dan tempo pembayaran. Sistem menampilkan status risiko pelanggan berdasarkan perbandingan piutang terhadap batas kreditnya.")
para("Pada tab pemasok, pengguna dapat menyimpan nama, nomor telepon, dan alamat pemasok. Data kontak ini kemudian digunakan dalam transaksi penjualan dan pembelian, sehingga setiap transaksi dapat dikaitkan dengan pelanggan atau pemasok yang bersangkutan.")

H3("4.2.5  Pembelian")
para("Halaman pembelian digunakan untuk mencatat barang yang masuk dari pemasok. Pengguna memilih pemasok, menambahkan barang beserta jumlah dan harga pokoknya, lalu menentukan metode pembayaran. Nomor nota pembelian dibuat secara otomatis untuk memudahkan penelusuran.")
para("Ketika pembelian disimpan, stok barang bertambah secara otomatis dan harga pokok diperbarui mengikuti rata-rata nilai persediaan. Apabila pembelian dilakukan secara kredit, sisa pembayaran dicatat sebagai hutang kepada pemasok dengan tanggal jatuh tempo yang ditetapkan.")

H3("4.2.6  Penjualan")
para("Halaman penjualan digunakan untuk melayani transaksi penjualan kepada pelanggan. Pengguna menambahkan barang beserta jumlahnya, dan harga jual akan terisi otomatis dari data barang namun tetap dapat disesuaikan. Sistem menghitung subtotal, potongan, biaya pengiriman, dan total secara otomatis.")
para("Sebelum disimpan, sistem memeriksa kecukupan stok dan, untuk penjualan kredit, memeriksa batas kredit pelanggan. Setelah transaksi tersimpan, stok barang berkurang secara otomatis. Bila pembayaran belum lunas, sisanya dicatat sebagai piutang. Nota penjualan dapat dicetak atau disimpan dalam format PDF.")

H3("4.2.7  Manajemen Stok, Penyesuaian, dan Mutasi")
para("Manajemen stok berjalan secara otomatis seiring transaksi pembelian dan penjualan. Untuk menjaga keakuratan, aplikasi menyediakan fitur penyesuaian stok yang digunakan ketika terjadi kerusakan, kehilangan, atau perbedaan hasil penghitungan fisik. Pengguna memasukkan perubahan jumlah beserta alasannya.")
para("Halaman mutasi stok menyajikan pergerakan stok setiap barang dalam suatu periode, meliputi stok awal, barang masuk, barang keluar, penyesuaian, dan stok akhir. Laporan ini membantu pengguna menelusuri perubahan stok serta memastikan jumlah barang di sistem sesuai dengan kondisi sebenarnya.")

H3("4.2.8  Pengeluaran")
para("Halaman pengeluaran digunakan untuk mencatat biaya operasional usaha di luar pembelian barang. Pengguna memilih kategori pengeluaran seperti gaji, sewa, transportasi, listrik dan air, atau kategori lainnya, kemudian memasukkan tanggal dan jumlahnya.")
para("Pencatatan pengeluaran yang tertib penting karena turut menentukan perhitungan laba bersih. Seluruh pengeluaran yang dicatat akan diperhitungkan dalam laporan laba rugi sebagai beban operasional pada periode yang bersangkutan.")

H3("4.2.9  Piutang dan Hutang")
para("Halaman piutang dan hutang membantu pengguna memantau kewajiban yang belum terselesaikan. Bagian piutang menampilkan penjualan kredit yang belum lunas beserta tanggal jatuh temponya, sedangkan bagian hutang menampilkan pembelian kredit yang masih harus dibayar kepada pemasok.")
para("Pada masing-masing daftar, pengguna dapat mencatat pembayaran sebagian maupun pelunasan penuh. Transaksi yang telah melewati tanggal jatuh tempo ditandai secara khusus agar pengguna dapat segera menindaklanjuti, baik dengan menagih pelanggan maupun menjadwalkan pembayaran kepada pemasok.")

H3("4.2.10  Pembukuan dan Laporan")
para("Fitur pembukuan dan laporan merupakan inti dari aplikasi ini. Halaman laporan menyajikan laporan laba rugi yang memuat pendapatan, harga pokok penjualan, laba kotor, beban operasional, dan laba bersih. Pengguna dapat menyaring laporan berdasarkan rentang tanggal sesuai kebutuhan.")
para("Selain laba rugi, halaman laporan menyajikan arus kas, posisi keuangan sederhana, laba per kategori barang, serta analisis umur piutang dan umur hutang yang dikelompokkan menurut rentang waktu. Penyajian dalam bentuk tabel dan grafik membuat informasi keuangan mudah dipahami dan dimanfaatkan untuk pengambilan keputusan.")

H3("4.2.11  Export Excel")
para("Aplikasi menyediakan fitur ekspor data ke berkas Excel menggunakan pustaka XLSX. Fitur ini tersedia pada berbagai halaman seperti data barang, penjualan, pembelian, mutasi stok, dan laporan. Data yang diekspor disusun rapi lengkap dengan baris total sehingga dapat diolah lebih lanjut di luar aplikasi.")
para("Kemampuan ekspor ini bermanfaat ketika pengguna perlu mengarsipkan data, menyusun laporan tambahan, atau membagikan informasi kepada pihak lain. Laporan keuangan bahkan dapat diekspor dalam beberapa lembar sekaligus agar pengguna memperoleh berkas yang lengkap.")

H3("4.2.12  Cetak Nota PDF")
para("Fitur pencetakan nota memanfaatkan pustaka jsPDF untuk menghasilkan dokumen dalam format PDF. Nota memuat informasi lengkap seperti nomor nota, tanggal, daftar barang, jumlah, harga, dan total pembayaran, sehingga dapat diberikan kepada pelanggan sebagai bukti transaksi yang resmi dan rapi.")
para("Selain nota penjualan, berbagai daftar dan laporan juga dapat dicetak maupun disimpan dalam format PDF. Kemampuan ini menggantikan peran nota kertas yang mudah hilang dan memberikan kesan profesional pada pelayanan toko.")

H3("4.2.13  Multi Bahasa")
para("Aplikasi mendukung dua bahasa, yaitu Bahasa Indonesia dan Bahasa Inggris, melalui pustaka i18next. Pengguna dapat mengganti bahasa melalui halaman pengaturan, dan seluruh teks pada menu, tombol, label, serta pesan akan menyesuaikan secara otomatis.")
para("Dukungan dua bahasa memperluas jangkauan pengguna aplikasi. Fitur ini juga memudahkan apabila aplikasi digunakan oleh pengguna dengan preferensi bahasa yang berbeda, tanpa perlu mengubah pengaturan pada tingkat sistem.")

H3("4.2.14  Pengaturan")
para("Halaman pengaturan menyediakan pengelolaan profil, pemilihan bahasa, serta pengelolaan data. Pada bagian pengelolaan data, pengguna dapat menyetel ulang data tertentu, seperti hanya data transaksi atau seluruh data, dengan mekanisme konfirmasi yang ketat untuk mencegah penghapusan yang tidak disengaja.")
para("Mekanisme konfirmasi mewajibkan pengguna mengetikkan frasa tertentu sebelum penyetelan ulang dijalankan. Pendekatan ini menjaga keamanan data sekaligus memberikan keleluasaan bagi pengguna yang ingin memulai pencatatan dari awal.")

H2("4.3  Cara Kerja Aplikasi Secara Umum")
para("Secara umum, aplikasi Bangun Kelola bekerja dengan menghubungkan antarmuka yang digunakan pengguna dengan basis data yang menyimpan seluruh informasi. Ketika pengguna memasukkan suatu transaksi, data tersebut dikirim ke server untuk disimpan, kemudian sistem memperbarui informasi terkait seperti stok dan saldo secara otomatis.")
para("Pembaruan data yang terjadi di basis data langsung tercermin pada tampilan aplikasi. Sebagai contoh, ketika sebuah penjualan dicatat, stok barang berkurang, piutang bertambah apabila pembayaran belum lunas, dan laporan keuangan ikut menyesuaikan. Seluruh perubahan ini berlangsung tanpa pengguna perlu menghitung ulang secara manual.")
para("Mekanisme keamanan memastikan setiap pengguna hanya dapat melihat dan mengubah data miliknya sendiri. Dengan demikian, meskipun aplikasi digunakan oleh banyak toko sekaligus, data setiap toko tetap terpisah dan terjaga kerahasiaannya.")

H2("4.4  Peran Masing-Masing Teknologi")
para("Setiap teknologi yang digunakan memiliki peran tersendiri dalam membentuk aplikasi yang utuh. Rangkuman peran teknologi tersebut disajikan pada tabel berikut.")
table(
    ["Teknologi", "Peran dalam Aplikasi"],
    [
        ["TypeScript", "Bahasa pemrograman utama yang menjaga keandalan kode."],
        ["React 19", "Membangun antarmuka pengguna berbasis komponen."],
        ["TanStack Start & Router", "Mengatur kerangka aplikasi dan perpindahan antarhalaman."],
        ["React Query", "Mengelola pengambilan dan pembaruan data dari server."],
        ["Tailwind CSS", "Menata tampilan secara cepat dan konsisten."],
        ["Radix UI & Lucide React", "Menyediakan komponen antarmuka dan ikon."],
        ["Recharts", "Menyajikan data dalam bentuk grafik."],
        ["React Hook Form & Zod", "Mengelola dan memvalidasi formulir."],
        ["Supabase", "Menangani autentikasi dan layanan basis data."],
        ["PostgreSQL", "Menyimpan seluruh data secara terstruktur dan aman."],
        ["jsPDF", "Membuat nota dan dokumen dalam format PDF."],
        ["XLSX", "Mengekspor data ke berkas Excel."],
        ["i18next", "Menyediakan dukungan dua bahasa."],
        ["date-fns", "Mengolah dan memformat tanggal."],
        ["Vercel", "Menerapkan aplikasi agar dapat diakses daring."],
        ["Bun", "Mengelola paket dan proses pengembangan."],
    ],
    widths=[2.1, 4.6],
)

H2("4.5  Alasan Pemilihan Teknologi")
para("Pemilihan teknologi pada aplikasi Bangun Kelola didasarkan pada pertimbangan kemudahan pengembangan, keandalan, serta kesesuaian dengan kebutuhan aplikasi web modern. React dipilih karena populer, didukung komunitas yang besar, dan unggul dalam membangun antarmuka yang responsif dengan pendekatan komponen yang mudah dikelola.")
para("TypeScript dipilih untuk meningkatkan keandalan, khususnya karena aplikasi banyak menangani data keuangan yang menuntut ketelitian. Supabase dengan basis data PostgreSQL dipilih karena menyediakan layanan backend yang lengkap dan aman secara siap pakai, sehingga mempercepat pengembangan tanpa mengorbankan keamanan data.")
para("Tailwind CSS dipilih karena mempercepat penataan tampilan sekaligus menjaga konsistensi desain. Adapun pustaka pendukung seperti jsPDF, XLSX, dan i18next dipilih untuk memenuhi kebutuhan spesifik berupa pencetakan nota, ekspor data, dan dukungan banyak bahasa. Penerapan melalui Vercel dan pengelolaan dengan Bun dipilih karena prosesnya praktis dan efisien.")

H2("4.6  Hasil Implementasi")
para("Hasil implementasi menunjukkan bahwa seluruh fitur yang direncanakan berhasil diwujudkan dan dapat berfungsi sebagaimana mestinya. Aplikasi Bangun Kelola mampu menangani proses pencatatan penjualan, pembelian, manajemen stok, pembukuan, laporan keuangan, dan pencetakan nota dalam satu sistem yang terintegrasi.")
para("Pengguna dapat menjalankan seluruh alur usaha mulai dari penyiapan data master, pencatatan transaksi harian, hingga peninjauan laporan keuangan. Stok barang dan informasi keuangan diperbarui secara otomatis setiap kali terjadi transaksi, sehingga data yang disajikan selalu mutakhir dan akurat.")
para("Antarmuka aplikasi tampil rapi dan responsif baik pada komputer maupun perangkat bergerak. Dengan demikian, aplikasi dapat digunakan dalam situasi kerja yang beragam, baik di meja kasir maupun saat memeriksa kondisi stok di gudang.")

H2("4.7  Pengujian Aplikasi")
para("Pengujian dilakukan menggunakan metode kotak hitam (black-box), yaitu pengujian yang berfokus pada kesesuaian keluaran aplikasi terhadap masukan tertentu tanpa memeriksa kode program di dalamnya. Pengujian ini bertujuan memastikan setiap fitur berjalan sesuai harapan. Hasil pengujian disajikan pada tabel berikut.")
table(
    ["No", "Skenario Pengujian", "Hasil yang Diharapkan", "Hasil"],
    [
        ["1", "Pendaftaran akun baru dengan data lengkap", "Akun terbuat dan pengguna masuk ke aplikasi", "Sesuai"],
        ["2", "Masuk dengan email dan kata sandi yang benar", "Pengguna berhasil masuk ke halaman utama", "Sesuai"],
        ["3", "Masuk dengan kata sandi salah", "Sistem menolak dan menampilkan pesan kesalahan", "Sesuai"],
        ["4", "Menambah barang baru beserta stok awal", "Barang tersimpan dan tampil pada daftar", "Sesuai"],
        ["5", "Mencatat pembelian barang", "Stok barang bertambah secara otomatis", "Sesuai"],
        ["6", "Mencatat penjualan barang", "Stok barang berkurang secara otomatis", "Sesuai"],
        ["7", "Menjual barang melebihi stok tersedia", "Sistem menolak transaksi", "Sesuai"],
        ["8", "Penjualan kredit melebihi batas kredit pelanggan", "Sistem menolak transaksi", "Sesuai"],
        ["9", "Mencatat pembayaran piutang", "Saldo piutang berkurang sesuai pembayaran", "Sesuai"],
        ["10", "Mencatat pengeluaran operasional", "Pengeluaran tercatat dan masuk laporan laba rugi", "Sesuai"],
        ["11", "Melihat laporan laba rugi pada periode tertentu", "Laporan tampil sesuai rentang tanggal", "Sesuai"],
        ["12", "Mencetak nota penjualan ke format PDF", "Nota PDF berhasil dibuat", "Sesuai"],
        ["13", "Mengekspor data ke berkas Excel", "Berkas Excel berhasil diunduh", "Sesuai"],
        ["14", "Mengganti bahasa aplikasi", "Seluruh teks berubah sesuai bahasa terpilih", "Sesuai"],
    ],
    widths=[0.4, 2.6, 2.7, 1.0],
)
para("Berdasarkan hasil pengujian, seluruh skenario memberikan hasil yang sesuai dengan harapan. Hal ini menunjukkan bahwa fitur-fitur utama aplikasi Bangun Kelola telah berfungsi dengan baik dan siap digunakan untuk mendukung kegiatan operasional toko bangunan.")

H2("4.8  Analisis Kelebihan Aplikasi")
para("Aplikasi Bangun Kelola memiliki sejumlah kelebihan yang membedakannya dari pencatatan manual maupun aplikasi umum. Kelebihan-kelebihan tersebut diuraikan sebagai berikut.")
bullet("Seluruh proses usaha tercatat dalam satu sistem terintegrasi sehingga data saling terhubung dan konsisten.")
bullet("Stok dan laporan keuangan diperbarui secara otomatis setiap terjadi transaksi, sehingga mengurangi kesalahan perhitungan.")
bullet("Istilah dan alur kerja disesuaikan dengan kebutuhan toko bangunan sehingga mudah dipahami pelaku usaha.")
bullet("Aplikasi berbasis web yang dapat diakses melalui peramban tanpa pemasangan khusus dan tampil responsif di berbagai perangkat.")
bullet("Data setiap toko terpisah dan aman karena diterapkan mekanisme keamanan tingkat baris.")
bullet("Tersedia fitur pencetakan nota PDF, ekspor Excel, dan dukungan dua bahasa yang menambah nilai guna aplikasi.")

H2("4.9  Analisis Kekurangan Aplikasi")
para("Di samping kelebihannya, aplikasi ini masih memiliki beberapa keterbatasan yang dapat menjadi bahan pengembangan pada masa mendatang. Keterbatasan tersebut antara lain sebagai berikut.")
bullet("Aplikasi memerlukan koneksi internet untuk dapat digunakan karena bersifat daring.")
bullet("Pembagian peran pengguna belum dipisahkan secara teknis, sehingga akses pemilik, admin, dan kasir masih berada dalam satu akun.")
bullet("Laporan keuangan yang disajikan masih bersifat praktis dan belum sepenuhnya mengikuti standar akuntansi keuangan secara penuh.")
bullet("Aplikasi belum dilengkapi fitur pemindaian kode batang untuk mempercepat pencarian barang.")
bullet("Belum tersedia fitur cadangan data otomatis yang dapat diatur oleh pengguna secara mandiri.")

page_break()

# ============================================================
# BAB V PENUTUP
# ============================================================
H1("BAB V  PENUTUP")

H2("5.1  Kesimpulan")
para("Berdasarkan uraian pada bab-bab sebelumnya, dapat ditarik beberapa kesimpulan dari pengembangan aplikasi Bangun Kelola sebagai berikut.")
bullet("Pengelolaan pembukuan pada toko bangunan yang masih dilakukan secara manual menimbulkan berbagai kendala, seperti catatan yang mudah hilang, perhitungan yang rentan keliru, serta sulitnya mengetahui kondisi keuangan secara tepat waktu.", num=True)
bullet("Aplikasi pembukuan dan akuntansi berbasis web berhasil dirancang dengan menyesuaikan diri terhadap proses bisnis toko bangunan, mulai dari arus pembelian, penjualan, hingga penyusunan laporan keuangan.", num=True)
bullet("Seluruh fitur yang direncanakan, meliputi pencatatan penjualan, pembelian, manajemen stok, pembukuan, laporan keuangan, dan pencetakan nota, berhasil diimplementasikan dalam satu sistem yang terintegrasi.", num=True)
bullet("Aplikasi Bangun Kelola terbukti dapat membantu pemilik toko, admin, dan kasir dalam mengelola usaha secara lebih efektif, akurat, dan cepat dibandingkan pencatatan manual, sebagaimana ditunjukkan oleh hasil pengujian yang seluruhnya sesuai harapan.", num=True)
para("Dengan demikian, tujuan pengembangan aplikasi telah tercapai. Bangun Kelola dapat menjadi solusi digitalisasi pembukuan yang relevan bagi pelaku usaha toko bangunan dalam meningkatkan kualitas pengelolaan administrasi dan keuangannya.")

H2("5.2  Saran Pengembangan")
para("Untuk meningkatkan kebermanfaatan dan kelengkapan aplikasi pada masa mendatang, disampaikan beberapa saran pengembangan sebagai berikut.")
bullet("Menambahkan pembagian peran pengguna secara teknis agar hak akses pemilik, admin, dan kasir dapat dipisahkan sesuai tanggung jawab masing-masing.", num=True)
bullet("Melengkapi aplikasi dengan kemampuan beroperasi secara luring terbatas sehingga tetap dapat digunakan ketika koneksi internet terganggu.", num=True)
bullet("Menyediakan fitur pemindaian kode batang untuk mempercepat pencarian dan pencatatan barang pada saat transaksi.", num=True)
bullet("Mengembangkan penyajian laporan keuangan agar lebih mendekati standar akuntansi keuangan serta menambahkan fitur cadangan data otomatis.", num=True)
bullet("Menambahkan fitur notifikasi, misalnya pengingat stok menipis dan pengingat piutang yang mendekati jatuh tempo, untuk membantu pengguna mengambil tindakan secara tepat waktu.", num=True)

page_break()

# ============================================================
# DAFTAR PUSTAKA
# ============================================================
H1("DAFTAR PUSTAKA")
para("Daftar pustaka berikut disusun mengikuti gaya penulisan American Psychological Association edisi ketujuh. Tahun, judul, dan penerbit pada rujukan buku dapat disesuaikan dengan sumber yang benar-benar digunakan oleh penyusun.", indent=False)


def ref(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(8)
    p.add_run(text)


ref("Hall, J. A. (2019). Accounting information systems (10th ed.). Cengage Learning.")
ref("Kieso, D. E., Weygandt, J. J., & Warfield, T. D. (2020). Intermediate accounting (17th ed.). John Wiley & Sons.")
ref("Laudon, K. C., & Laudon, J. P. (2021). Management information systems: Managing the digital firm (17th ed.). Pearson Education.")
ref("Mulyadi. (2016). Sistem akuntansi (Edisi keempat). Salemba Empat.")
ref("Romney, M. B., & Steinbart, P. J. (2021). Accounting information systems (15th ed.). Pearson Education.")
ref("Sommerville, I. (2016). Software engineering (10th ed.). Pearson Education.")
ref("PostgreSQL Global Development Group. (2024). PostgreSQL documentation. Diakses dari https://www.postgresql.org/docs/")
ref("Supabase. (2024). Supabase documentation. Diakses dari https://supabase.com/docs")
ref("Meta Open Source. (2024). React documentation. Diakses dari https://react.dev/")
ref("Tailwind Labs. (2024). Tailwind CSS documentation. Diakses dari https://tailwindcss.com/docs")
ref("Microsoft. (2024). TypeScript documentation. Diakses dari https://www.typescriptlang.org/docs/")
ref("TanStack. (2024). TanStack Router and Start documentation. Diakses dari https://tanstack.com/")

page_break()

# ============================================================
# LAMPIRAN
# ============================================================
H1("LAMPIRAN")
para("Bagian lampiran memuat dokumen pendukung yang melengkapi laporan. Berikut adalah daftar lampiran yang disarankan untuk disertakan. Tangkapan layar dan diagram dapat ditambahkan pada bagian ini sesuai dengan kondisi aplikasi yang dikembangkan.", indent=False)
items = [
    ("Lampiran 1", "Tangkapan layar halaman Dasbor."),
    ("Lampiran 2", "Tangkapan layar halaman Data Barang."),
    ("Lampiran 3", "Tangkapan layar halaman Penjualan."),
    ("Lampiran 4", "Tangkapan layar halaman Pembelian."),
    ("Lampiran 5", "Tangkapan layar halaman Laporan Keuangan."),
    ("Lampiran 6", "Tangkapan layar hasil Cetak Nota dalam format PDF."),
    ("Lampiran 7", "Struktur Basis Data aplikasi."),
    ("Lampiran 8", "Entity Relationship Diagram (ERD)."),
    ("Lampiran 9", "Use Case Diagram."),
    ("Lampiran 10", "Activity Diagram."),
    ("Lampiran 11", "Manual Penggunaan Aplikasi."),
]
table(
    ["Lampiran", "Keterangan"],
    [[a, b] for a, b in items],
    widths=[1.3, 5.4],
)
for a, b in items:
    figure_placeholder(a + ": " + b)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ============================================================
# SIMPAN
# ============================================================
out = "Laporan_Bangun_Kelola.docx"
doc.save(out)
print("Dokumen berhasil dibuat:", out)
